import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter
from typing import List, Dict, Any

class ExcelExporter:
    @staticmethod
    def export_table_to_excel(file_path: str, headers: List[str], rows: List[List[Any]], title: str = "گزارش"):
        """
        Exports headers and rows to a styled .xlsx file with RTL direction.
        """
        wb = openpyxl.Workbook()
        ws = wb.active

        # Clean title for worksheet name (openpyxl disallows invalid characters like : \ / ? * [ ])
        clean_title = "".join(c for c in title if c not in r":\/?*[]")[:30].strip() or "گزارش"
        ws.title = clean_title
        ws.views.sheetView[0].rightToLeft = True

        # Styles
        title_font = Font(name="Tahoma", size=14, bold=True, color="2C3E50")
        header_font = Font(name="Tahoma", size=11, bold=True, color="FFFFFF")
        header_fill = PatternFill(start_color="2980B9", end_color="2980B9", fill_type="solid")

        data_font = Font(name="Tahoma", size=10)
        alt_fill = PatternFill(start_color="F2F4F4", end_color="F2F4F4", fill_type="solid")

        thin_border = Border(
            left=Side(style='thin', color='D0D7DE'),
            right=Side(style='thin', color='D0D7DE'),
            top=Side(style='thin', color='D0D7DE'),
            bottom=Side(style='thin', color='D0D7DE')
        )

        align_center = Alignment(horizontal='center', vertical='center')
        align_right = Alignment(horizontal='right', vertical='center')

        # Title row
        ws.merge_cells(start_row=1, start_column=1, end_row=1, end_column=len(headers))
        cell_title = ws.cell(row=1, column=1, value=title)
        cell_title.font = title_font
        cell_title.alignment = align_center

        # Header row
        for col_idx, header in enumerate(headers, 1):
            cell = ws.cell(row=3, column=col_idx, value=header)
            cell.font = header_font
            cell.fill = header_fill
            cell.alignment = align_center
            cell.border = thin_border

        # Data rows
        for row_idx, row_data in enumerate(rows, 4):
            is_even = (row_idx % 2 == 0)
            for col_idx, val in enumerate(row_data, 1):
                cell = ws.cell(row=row_idx, column=col_idx, value=val)
                cell.font = data_font
                cell.border = thin_border
                cell.alignment = align_right
                if is_even:
                    cell.fill = alt_fill

        # Auto-fit columns safely
        for col_idx in range(1, len(headers) + 1):
            max_len = 0
            col_letter = get_column_letter(col_idx)
            for row_idx in range(3, len(rows) + 4):
                val_str = str(ws.cell(row=row_idx, column=col_idx).value or '')
                if len(val_str) > max_len:
                    max_len = len(val_str)
            ws.column_dimensions[col_letter].width = min(max(max_len + 6, 12), 60)

        wb.save(file_path)
